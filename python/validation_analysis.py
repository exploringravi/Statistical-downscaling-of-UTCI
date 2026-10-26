"""Validation analysis for the R1 revision of remotesensing-4629248.

Reads the observation and model pairs in "Validation data/" and writes, to
"validation_output/":
  validation_metrics.csv      Tables 6 and 7, Supplementary S1 and S2
  skill_scores.csv            Supplementary S6
  utci_error_nicosia.csv      Supplementary S10 (Nicosia stations)
  Figure10_Perugia.png/pdf, Figure11_Limassol.png/pdf,
  Figure12_Nicosia_P1_P2.png/pdf, Figure14_station_tests.png

Run from the repository root:
  python python/validation_analysis.py

Requirements: numpy, scipy, openpyxl, matplotlib, thermofeel
"""
import csv
import datetime as dt
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy.stats import pearsonr, spearmanr

DATA = "Validation data"
OUT = "validation_output"
VAR = {"Ta": ("Air temperature", "°C"), "RH": ("Relative humidity", "%"), "WS": ("Wind speed", "m/s")}

# ---------------------------------------------------------------- loading
def _parse_time(s):
    for fmt in ("%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(s, fmt)
        except (TypeError, ValueError):
            pass
    return None

def read_sheet(fname, sheet, time_sorted=False):
    """Return observed and modelled arrays from one variable sheet."""
    rows = list(openpyxl.load_workbook(os.path.join(DATA, fname), data_only=True)[sheet].iter_rows(values_only=True))
    hi = [k for k, r in enumerate(rows) if r[0] in ("Point_ID", "Point")][0]
    head = rows[hi]
    oc = [j for j, h in enumerate(head) if h and h.startswith("Observed")][0]
    mc = [j for j, h in enumerate(head) if h and h.startswith("Modelled")][0]
    data = [r for r in rows[hi + 1:] if isinstance(r[0], (int, float))]
    if time_sorted:
        tj = [j for j, h in enumerate(head) if h and h.startswith("Timestamp")][0]
        data = sorted(data, key=lambda r: _parse_time(r[tj]))
    return np.array([r[oc] for r in data], float), np.array([r[mc] for r in data], float)

def read_nicosia(station):
    rows = list(openpyxl.load_workbook(os.path.join(DATA, "Nicosia_P1_P2_2026.xlsx"), data_only=True)["All_Raw_Data"].iter_rows(values_only=True))[3:]
    rows = [r for r in rows if r[0] == station]
    cols = {"Ta": 2, "RH": 4, "WS": 6}
    return {v: (np.array([r[c] for r in rows], float), np.array([r[c + 1] for r in rows], float)) for v, c in cols.items()}

def load_all():
    """Return {(city, site, variable): (observed, modelled, bootstrap block length)}."""
    d = {}
    perugia = [
        ("Pian di Massiano 2021", "Perugia_Pian_di_Massiano_2021.xlsx", ["a_Ta_Air_Temperature", "b_RH_Relative_Humidity", "c_WS_Wind_Speed"]),
        ("UniStraPg 2022", "Perugia_UniStraPg_2022.xlsx", ["d_Ta_Air_Temperature", "e_RH_Relative_Humidity", "f_WS_Wind_Speed"]),
        ("Cortonese 2023", "Perugia_Cortonese_2023.xlsx", ["g_Ta_Air_Temperature", "h_RH_Relative_Humidity", None]),
    ]
    for site, f, sheets in perugia:
        for v, sh in zip(["Ta", "RH", "WS"], sheets):
            if sh:
                o, m = read_sheet(f, sh)
                d[("Perugia", site, v)] = (o, m, 1)  # no time order: individual resampling
    for site, pre in [("Limassol Route 1", "Ayia"), ("Limassol Route 2 (Episkopi)", "Episkopi")]:
        for v in VAR:
            o, m = read_sheet("Limassol_Route1_Route2_2025.xlsx", f"{pre}_{v}", time_sorted=True)
            d[("Limassol", site, v)] = (o, m, max(1, round(0.1 * len(o))))
    for st in ["P1", "P2"]:
        for v, (o, m) in read_nicosia(st).items():
            d[("Nicosia", st, v)] = (o, m, 2)  # 2 h blocks
    return d

# ---------------------------------------------------------------- statistics
def resample_index(n, block, rng):
    if block <= 1:
        return rng.integers(0, n, n)
    nb = int(np.ceil(n / block))
    starts = rng.integers(0, n - block + 1, nb)
    return np.concatenate([np.arange(s, s + block) for s in starts])[:n]

def rmse_ci(o, m, block, rng, b=10000):
    r = m - o
    vals = [np.sqrt((r[resample_index(len(r), block, rng)] ** 2).mean()) for _ in range(b)]
    return np.percentile(vals, [2.5, 97.5])

def skill_ci(o, m, block, rng, b=5000):
    vals = []
    for _ in range(b):
        i = resample_index(len(o), block, rng)
        vals.append(1 - (m[i] - o[i]).std() / o[i].std())
    return np.percentile(vals, [2.5, 97.5])

def willmott_d(o, m):
    ob = o.mean()
    return 1 - ((m - o) ** 2).sum() / ((abs(m - ob) + abs(o - ob)) ** 2).sum()

def metrics(o, m):
    r = m - o
    slope, icpt = np.polyfit(o, m, 1)
    return dict(n=len(o), rmse=np.sqrt((r ** 2).mean()), mae=abs(r).mean(), bias=r.mean(),
                medae=np.median(abs(r)), offset_removed_rmse=r.std(), obs_sd=o.std(),
                d=willmott_d(o, m), r=pearsonr(o, m)[0], rho=spearmanr(o, m)[0],
                slope=slope, intercept=icpt)

# ---------------------------------------------------------------- UTCI (S10)
def utci_c(ta, rh, ws15, mrt_offset):
    import thermofeel as tf
    tk = ta + 273.15
    e = rh / 100 * tf.calculate_saturation_vapour_pressure(tk)
    va10 = np.maximum(ws15 / 0.58, 0.5)  # base pedestrian factor; 0.5 m/s polynomial floor
    return tf.calculate_utci(tk, va10, tk + mrt_offset, ehPa=e) - 273.15

# ---------------------------------------------------------------- figures
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": "#444",
                     "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white"})
HEADERS = {
    ("Perugia", "Pian di Massiano 2021"): "Pian di Massiano-Barton Park  |  15 June 2021  |  18:00 to 19:40 CEST",
    ("Perugia", "UniStraPg 2022"): "UniStra-Viale Carlo Manuali  |  19 July 2022  |  12:40 to 14:00 CEST",
    ("Perugia", "Cortonese 2023"): "Cortonese-Madonna Alta  |  8 August 2023  |  12:40 to 13:00 CEST",
    ("Limassol", "Limassol Route 1"): "Ayia Napa Bus Stop to Limassol Castle route  |  21 May 2025  |  19:55 to 20:45 EEST",
    ("Limassol", "Limassol Route 2 (Episkopi)"): "Episkopi area  |  21 May 2025  |  18:00 to 19:00 EEST",
    ("Nicosia", "P1"): "P1, shaded courtyard, Agios Andreas Primary School, Nicosia  |  22 May 2026  |  24 hourly values",
    ("Nicosia", "P2"): "P2, open parking area, Agios Andreas Primary School, Nicosia  |  22 May 2026  |  24 hourly values",
}

def scatter_panel(ax, o, m, v, label):
    name, u = VAR[v]
    k = metrics(o, m)
    lo, hi = min(o.min(), m.min()), max(o.max(), m.max())
    pad = (hi - lo) * 0.08
    lo, hi = lo - pad, hi + pad
    ax.plot([lo, hi], [lo, hi], color="#555", lw=1.1, label="1:1")
    ax.scatter(o, m, s=22, color="#2a6fb0", alpha=0.75, edgecolor="white", linewidth=0.6, zorder=3)
    xx = np.array([lo, hi])
    ax.plot(xx, k["slope"] * xx + k["intercept"], color="#1a1a1a", lw=1.4, ls="--", label="Fit")
    ax.set_xlim(lo, hi); ax.set_ylim(lo, hi); ax.set_aspect("equal")
    ax.set_xlabel(f"Observed {name} ({u})"); ax.set_ylabel(f"Modelled {name} ({u})")
    ax.grid(color="#e6e6e6", lw=0.6); ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.set_title(f"({label}) {name}", loc="left", fontsize=9.5, fontweight="bold")
    sign = "+" if k["intercept"] >= 0 else "-"
    txt = (f"n = {k['n']}\nRMSE = {k['rmse']:.2f} {u}\nMAE = {k['mae']:.2f} {u}\nBias = {k['bias']:+.2f} {u}\n"
           f"Willmott d = {k['d']:.2f}\ny = {k['slope']:.2f}x {sign} {abs(k['intercept']):.2f}")
    ax.text(0.03, 0.97, txt, transform=ax.transAxes, va="top", fontsize=7.6,
            bbox=dict(boxstyle="round,pad=0.35", fc="white", ec="#ccc", lw=0.6))

def scatter_figure(data, city, sites, title, fname):
    fig = plt.figure(figsize=(10.5, 3.85 * len(sites) + 0.55))
    sfs = fig.subfigures(len(sites) + 1, 1, height_ratios=[0.55 / 3.85] + [1] * len(sites))
    top = sfs[0]
    top.text(0.5, 0.45, title, ha="center", va="center", fontsize=12.5, fontweight="bold")
    n = 0
    for sf, site in zip(sfs[1:], sites):
        sf.suptitle(HEADERS[(city, site)], fontsize=10, fontweight="bold", y=0.97)
        axs = sf.subplots(1, 3)
        sf.subplots_adjust(left=0.07, right=0.98, top=0.84, bottom=0.15, wspace=0.38)
        for ax, v in zip(axs, ["Ta", "RH", "WS"]):
            key = (city, site, v)
            if key not in data:
                ax.axis("off"); continue
            o, m, _ = data[key]
            scatter_panel(ax, o, m, v, "abcdefghi"[n]); n += 1
        if site == sites[0]:
            h, l = axs[0].get_legend_handles_labels()
            top.legend(h, l, loc="center right", frameon=False, ncol=2, fontsize=8.5)
    for ext in ("png", "pdf"):
        fig.savefig(os.path.join(OUT, f"{fname}.{ext}"), dpi=300)
    plt.close(fig)

def figure14(nicosia):
    # DWD values (Köln district, 14:00) as reported in Section 4.4
    labels = ["Nicosia station P1", "Nicosia station P2", "DWD stations, 15 Jul", "DWD stations, 31 Dec"]
    fw = [nicosia["P1"], nicosia["P2"], 0.80, 0.65]
    era = [1.31, 1.82, 0.91, 0.54]  # nearest ERA5-Land value
    fig, ax = plt.subplots(figsize=(10, 4.2))
    y = np.arange(len(labels))[::-1]; h = 0.36
    ax.barh(y + h / 2, fw, h, color="#1b7f79", label="10 m framework", edgecolor="white")
    ax.barh(y - h / 2, era, h, color="#9b9a95", label="ERA5-Land (~9 km)", edgecolor="white")
    for yy, a, b in zip(y, fw, era):
        ax.text(a + 0.03, yy + h / 2, f"{a:.2f}", va="center", fontsize=11)
        ax.text(b + 0.03, yy - h / 2, f"{b:.2f}", va="center", fontsize=11, color="#555")
    ax.set_yticks(y); ax.set_yticklabels(labels); ax.set_xlim(0, 2.2)
    ax.set_xlabel("Raw air-temperature RMSE (°C), lower is better")
    ax.set_title("10 m framework against ERA5-Land in the station tests", fontweight="bold", fontsize=13)
    ax.grid(axis="x", color="#e6e6e6"); ax.set_axisbelow(True)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.legend(loc="lower right", frameon=False)
    fig.tight_layout(); fig.savefig(os.path.join(OUT, "Figure14_station_tests.png"), dpi=200)
    plt.close(fig)

# ---------------------------------------------------------------- main
def main():
    os.makedirs(OUT, exist_ok=True)
    data = load_all()

    rng = np.random.default_rng(42)
    with open(os.path.join(OUT, "validation_metrics.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["City", "Site", "Variable", "n", "Observed_SD", "RMSE", "MAE", "Bias", "RMSE_CI_low", "RMSE_CI_high",
                    "Offset_removed_RMSE", "MedAE", "Willmott_d", "Pearson_r", "Spearman_rho"])
        for (city, site, v), (o, m, block) in data.items():
            k = metrics(o, m); lo, hi = rmse_ci(o, m, block, rng)
            w.writerow([city, site, v, k["n"]] + [round(x, 3) for x in (k["obs_sd"], k["rmse"], k["mae"], k["bias"], lo, hi,
                       k["offset_removed_rmse"], k["medae"], k["d"], k["r"], k["rho"])])

    rng = np.random.default_rng(7)
    with open(os.path.join(OUT, "skill_scores.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["City", "Site", "Variable", "n", "sigma_obs", "RMSE_bc", "r", "SS", "SS_CI_low", "SS_CI_high"])
        for (city, site, v), (o, m, block) in data.items():
            k = metrics(o, m); lo, hi = skill_ci(o, m, block, rng)
            ss = 1 - k["offset_removed_rmse"] / k["obs_sd"]
            w.writerow([city, site, v, k["n"]] + [round(x, 3) for x in (k["obs_sd"], k["offset_removed_rmse"], k["r"], ss, lo, hi)])

    # S10: only Nicosia has Ta, RH and WS matched at the same point
    with open(os.path.join(OUT, "utci_error_nicosia.csv"), "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Station", "n", "UTCI_RMSE_MRT_offset_0C", "UTCI_RMSE_MRT_offset_20C", "Pct_within_2C_offset_0", "Pct_within_2C_offset_20"])
        for st in ("P1", "P2"):
            s = read_nicosia(st)
            row = []
            for off in (0, 20):
                uo = utci_c(s["Ta"][0], s["RH"][0], s["WS"][0], off)
                um = utci_c(s["Ta"][1], s["RH"][1], s["WS"][1], off)
                dlt = um - uo
                row.append((np.sqrt((dlt ** 2).mean()), (abs(dlt) <= 2).mean() * 100))
            w.writerow([st, len(s["Ta"][0]), round(row[0][0], 3), round(row[1][0], 3), round(row[0][1], 1), round(row[1][1], 1)])

    scatter_figure(data, "Perugia", ["Pian di Massiano 2021", "UniStraPg 2022", "Cortonese 2023"],
                   "Validation of downscaled meteorological variables in Perugia", "Figure10_Perugia")
    scatter_figure(data, "Limassol", ["Limassol Route 1", "Limassol Route 2 (Episkopi)"],
                   "Validation of downscaled meteorological variables in Limassol", "Figure11_Limassol")
    scatter_figure(data, "Nicosia", ["P1", "P2"], "Temporal validation at two fixed stations in Nicosia", "Figure12_Nicosia_P1_P2")
    nic = {st: metrics(*data[("Nicosia", st, "Ta")][:2])["rmse"] for st in ("P1", "P2")}
    figure14({k: round(v, 2) for k, v in nic.items()})
    print(f"Outputs written to {OUT}/")

if __name__ == "__main__":
    main()
